# -*- coding: utf-8 -*-
"""budget.py — 예산 관리 창.

'예산관리 기능.html'의 로직(요약 카드 / 세부항목 탭 / 예산 테이블 /
집행액 인라인 편집 / 완료 토글 / 세부항목·항목 관리)을 PyQt5로 재구현.
데이터는 ssnnote.db의 budget_categories, budget_items 테이블에 저장.
"""

import os
import sys
import json

import qtawesome as qta
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QPushButton,
    QLabel, QLineEdit, QScrollArea, QFrame, QTableWidget, QTableWidgetItem,
    QHeaderView, QDialog, QMessageBox, QComboBox, QAbstractItemView,
    QProgressBar, QFileDialog, QCheckBox, QButtonGroup, QTextEdit, QPlainTextEdit
)
from PyQt5.QtCore import Qt, QSize
from PyQt5.QtGui import QColor, QFont, QFontDatabase, QKeySequence

from db import (get_budget_categories, add_budget_category, rename_budget_category,
                delete_budget_category, get_budget_items, add_budget_item,
                update_budget_item, delete_budget_item, set_budget_item_spent,
                toggle_budget_item_completed, clear_budget_data)


# ── 번들 Pretendard 등록 ──────────────────────────────────────────
# Pretendard가 설치되지 않은 PC에서도 예산 창이 Pretendard로 보이도록
# assets의 정적 폰트(여러 굵기)를 앱 폰트로 등록한다. QSS의
# font-family: 'Pretendard' 가 이 등록된 패밀리를 사용한다.
_pretendard_loaded = False


def _base_path():
    if getattr(sys, 'frozen', False):
        return sys._MEIPASS
    return os.path.dirname(os.path.abspath(__file__))


def _load_pretendard():
    """assets 안의 Pretendard 정적 폰트(.ttf/.otf)를 모두 등록한다. 1회만."""
    global _pretendard_loaded
    if _pretendard_loaded:
        return
    _pretendard_loaded = True
    for dirpath, _dirs, files in os.walk(os.path.join(_base_path(), 'assets')):
        for fn in sorted(files):
            low = fn.lower()
            if low.startswith('pretendard') and low.endswith(('.ttf', '.otf')):
                QFontDatabase.addApplicationFont(os.path.join(dirpath, fn))


# ── 전체 QSS 테마 ─────────────────────────────────────────────────
STYLE = """
    QMainWindow { background-color: #f1f5f9; }
    QWidget {
        font-family: 'Pretendard', 'Malgun Gothic';
        font-size: 10.5pt;
        color: #0f172a;
    }
    QScrollArea { border: none; background: transparent; }
    QFrame#Card { background: white; border: 1px solid #e2e8f0; border-radius: 10px; }
    QFrame#Section { background: white; border: 1px solid #cbd5e1; border-radius: 10px; }
    QFrame#SectionHeader {
        background: #1e1b4b;
        border: none;
        border-top-left-radius: 9px;
        border-top-right-radius: 9px;
    }
    QFrame#Toolbar { background: white; border: 1px solid #e2e8f0; border-radius: 10px; }
    QLabel#CardTitle { color: #64748b; font-size: 9.5pt; font-weight: 700; }
    QLabel#CardSub { color: #94a3b8; font-size: 8.5pt; }
    QLabel#CardValueIndigo { color: #4338ca; font-size: 19pt; font-weight: 800; }
    QLabel#CardValueAmber { color: #d97706; font-size: 19pt; font-weight: 800; }
    QLabel#CardValueEmerald { color: #059669; font-size: 19pt; font-weight: 800; }
    QLabel#PercentBig { color: #1d4ed8; font-size: 13pt; font-weight: 800; }
    QLabel#SectionTitle { color: white; font-size: 13pt; font-weight: 700; }
    QLabel#Chip {
        background: #312e81;
        border: 1px solid #4338ca;
        border-radius: 8px;
        padding: 4px 10px;
        font-size: 9.5pt;
        color: #ffffff;
    }
    QLabel#EmptyBox {
        background: white;
        border: 1px solid #e2e8f0;
        border-radius: 10px;
        color: #94a3b8;
        padding: 30px;
    }
    QPushButton#TabActive {
        background: #312e81; color: white; border: none; border-radius: 8px;
        padding: 6px 16px; font-weight: 700;
    }
    QPushButton#TabBtn {
        background: #e2e8f0; color: #334155; border: none; border-radius: 8px;
        padding: 6px 16px; font-weight: 600;
    }
    QPushButton#TabBtn:hover { background: #cbd5e1; }
    QPushButton#Primary {
        background: #4f46e5; color: white; border: none; border-radius: 8px;
        padding: 7px 14px; font-weight: 700;
    }
    QPushButton#Primary:hover { background: #4338ca; }
    QPushButton#Primary:disabled { background: #c7d2fe; }
    QPushButton#Danger {
        background: #be123c; color: white; border: none; border-radius: 8px;
        padding: 7px 12px; font-weight: 700;
    }
    QPushButton#Danger:hover { background: #9f1239; }
    QPushButton#Dark {
        background: #312e81; color: white; border: 1px solid #4338ca;
        border-radius: 8px; padding: 7px 12px; font-weight: 700;
    }
    QPushButton#Dark:hover { background: #3730a3; }
    QPushButton#Ghost {
        background: white; color: #334155; border: 1px solid #cbd5e1;
        border-radius: 8px; padding: 6px 12px; font-weight: 600;
    }
    QPushButton#Ghost:hover { background: #f1f5f9; }
    QPushButton#IconBtn { background: transparent; border: none; border-radius: 6px; }
    QPushButton#IconBtn:hover { background: #e2e8f0; }
    QPushButton#CompleteBtn {
        background: #e2e8f0; color: #334155; border: none; border-radius: 6px;
        padding: 4px 10px; font-weight: 700; font-size: 9.5pt;
    }
    QPushButton#CompleteBtn:hover { background: #cbd5e1; }
    QPushButton#DoneBtn {
        background: #059669; color: white; border: none; border-radius: 6px;
        padding: 4px 10px; font-weight: 700; font-size: 9.5pt;
    }
    QPushButton#DoneBtn:hover { background: #047857; }
    QTableWidget { background: white; border: none; gridline-color: transparent; }
    QTableWidget::item {
        border-right: 1px solid #e2e8f0;
        border-bottom: 1px solid #e2e8f0;
        padding: 4px 8px;
    }
    QHeaderView::section {
        background: #e2e8f0; color: #334155; padding: 7px 8px; border: none;
        border-right: 1px solid #cbd5e1; border-bottom: 1px solid #cbd5e1;
        font-weight: 700;
    }
    QLineEdit {
        background: white; border: 1px solid #cbd5e1; border-radius: 6px;
        padding: 5px 8px;
    }
    QLineEdit:focus { border: 2px solid #6366f1; }
    QLineEdit:disabled { background: #e2e8f0; color: #64748b; border-color: transparent; }
    QLineEdit#SpentInput {
        background: #fffbeb; border: 2px solid #fcd34d; color: #78350f;
        font-weight: 700;
    }
    QLineEdit#SpentInput:focus { border: 2px solid #6366f1; background: white; }
    QLineEdit#SpentInput:disabled {
        background: #e2e8f0; border-color: transparent; color: #64748b;
    }
    QComboBox {
        background: white; border: 1px solid #cbd5e1; border-radius: 6px;
        padding: 5px 8px;
    }
    QCheckBox { font-size: 12.5pt; spacing: 8px; }
    QCheckBox:disabled { color: #94a3b8; }
    QProgressBar { background: #e2e8f0; border: none; border-radius: 4px; }
    QProgressBar::chunk { background: #2563eb; border-radius: 4px; }
    QProgressBar#SummaryBar::chunk { background: #f59e0b; border-radius: 4px; }
    QScrollBar:vertical { width: 10px; background: transparent; }
    QScrollBar::handle:vertical { background: #cbd5e1; border-radius: 5px; min-height: 24px; }
    QScrollBar::handle:vertical:hover { background: #94a3b8; }
    QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height: 0; }
"""

_STATE_FILE = os.path.join(os.environ.get('APPDATA', '.'), 'SSNnote', 'budget_state.json')

_ROW_H = 46
_HEADER_H = 38

# 열 폭: 넉넉할 때의 폭과, 좁아져도 이 아래로는 줄이지 않는 최소 폭.
# 최소 폭은 실제 내용(헤더 텍스트·금액·진행바/버튼 셀)에 맞춘 값이라
# 숫자가 '1,200,0…'처럼 잘리지 않는다. 합계 939px + 창 여백 ≈ 최소 창 폭 1000.
_COL_W = (160, 258, 130, 160, 130, 150, 118, 86)
_COL_MIN = (120, 136, 124, 149, 112, 142, 90, 66)

_BG_ZEBRA = QColor('#f8fafc')
_BG_WHITE = QColor('#ffffff')
_BG_DONE = QColor('#f1f5f9')
_BG_NEG = QColor('#ffe4e6')

_COLOR_NEG = QColor('#be123c')
_COLOR_DONE_TEXT = QColor('#64748b')
_COLOR_TEXT = QColor('#0f172a')
_COLOR_COST = QColor('#475569')


def _fmt(n):
    try:
        return f'{int(n):,}'
    except (TypeError, ValueError):
        return '0'


def _parse_amount(text):
    digits = ''.join(ch for ch in str(text) if ch.isdigit())
    return int(digits) if digits else 0


def _pct(spent, budget):
    return (spent / budget * 100) if budget > 0 else 0.0


def _attach_amount_format(edit):
    def on_change(text):
        digits = ''.join(ch for ch in text if ch.isdigit())
        formatted = f'{int(digits):,}' if digits else ''
        if formatted != text:
            edit.blockSignals(True)
            edit.setText(formatted)
            edit.blockSignals(False)
            edit.setCursorPosition(len(formatted))
    edit.textChanged.connect(on_change)


def _make_strike_font(font):
    f = QFont(font)
    f.setStrikeOut(True)
    return f


def _icon_btn(icon_name, tooltip, color='#64748b', size=16):
    btn = QPushButton()
    btn.setObjectName('IconBtn')
    btn.setIcon(qta.icon(icon_name, color=color))
    btn.setIconSize(QSize(size, size))
    btn.setFixedSize(28, 28)
    btn.setToolTip(tooltip)
    return btn


_EXCEL_EXTS = ('.xls', '.xlsx')


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _to_number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(round(value))
    text = str(value).strip().replace(',', '').replace(' ', '')
    negative = False
    if text.startswith('(') and text.endswith(')'):
        negative = True
        text = text[1:-1]
    if text in ('', '-', '–', '—'):
        return None
    try:
        num = int(round(float(text)))
    except ValueError:
        return None
    return -num if negative else num


def _read_excel_sheets(path):
    """엑셀 파일을 [(시트이름, rows)] 형태로 읽는다. (.xls: xlrd, .xlsx: openpyxl)"""
    ext = os.path.splitext(path)[1].lower()
    sheets = []
    if ext == '.xls':
        import xlrd
        wb = xlrd.open_workbook(path)
        for sh in wb.sheets():
            rows = [[sh.cell_value(r, c) for c in range(sh.ncols)]
                    for r in range(sh.nrows)]
            sheets.append((sh.name, rows))
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        try:
            for sh in wb.worksheets:
                rows = [list(row) for row in sh.iter_rows(values_only=True)]
                sheets.append((sh.title, rows))
        finally:
            wb.close()
    else:
        raise ValueError('지원하지 않는 파일 형식입니다. (.xls, .xlsx)')
    return sheets


def _parse_budget_sheet(rows):
    """사업관리카드(예산) 시트를 세부항목/항목 구조로 파싱. 헤더 없으면 None."""
    header_row = None
    col_cost = col_desc = col_budget = col_spent = None
    for r in range(min(len(rows), 30)):
        texts = [_cell_text(v).replace(' ', '') for v in rows[r]]
        budget_idx = next((i for i, t in enumerate(texts) if '예산현액' in t), None)
        desc_idx = next((i for i, t in enumerate(texts) if '산출내역' in t), None)
        if budget_idx is None or desc_idx is None:
            continue
        header_row = r
        col_budget = budget_idx
        col_desc = desc_idx
        col_cost = next((i for i, t in enumerate(texts)
                         if '원가통계비목' in t or '세부사업' in t), 0)
        # 하위 헤더 행(집행 세부 열)까지 합쳐서 집행 열을 찾는다
        merged = list(texts)
        for extra in rows[r + 1:r + 3]:
            for i in range(min(len(merged), len(extra))):
                merged[i] += _cell_text(extra[i]).replace(' ', '')
        col_spent = next((i for i, t in enumerate(merged) if '원인행위' in t), None)
        if col_spent is None:
            col_spent = next((i for i, t in enumerate(merged) if '지출결의' in t), None)
        if col_spent is None:
            col_spent = next((i for i, t in enumerate(merged) if '지급액' in t), None)
        break
    if header_row is None:
        return None

    def cell(row, idx):
        return row[idx] if idx is not None and idx < len(row) else None

    categories = []
    current = None
    for row in rows[header_row + 1:]:
        cost = _cell_text(cell(row, col_cost))
        desc = _cell_text(cell(row, col_desc))
        budget = _to_number(cell(row, col_budget))
        spent = _to_number(cell(row, col_spent)) or 0
        if not cost and not desc and budget is None:
            continue
        if any(key in cost for key in ('합계', '총계', '소계')):
            continue
        if not desc and budget is not None:
            # 세부항목(카테고리) 소계 행
            current = {'name': cost or '미분류', 'items': []}
            categories.append(current)
            continue
        if desc:
            if current is None:
                current = {'name': '미분류', 'items': []}
                categories.append(current)
            current['items'].append({
                'cost_type': cost or '일반',
                'description': desc,
                'budget': budget or 0,
                'spent': spent,
            })
    return categories


def parse_budget_excel(path):
    """엑셀 파일 → [{'name': 세부항목, 'items': [...]}]. 실패 시 ValueError."""
    try:
        sheets = _read_excel_sheets(path)
    except ImportError as e:
        raise ValueError(f'엑셀 읽기 모듈을 불러오지 못했습니다: {e}')
    for _name, rows in sheets:
        if not rows:
            continue
        categories = _parse_budget_sheet(rows)
        if categories:
            return categories
    raise ValueError("'예산현액'/'산출내역' 헤더를 찾지 못했습니다.")


def _parse_tsv(text):
    """엑셀 클립보드 텍스트(탭 구분, 따옴표 셀 처리)를 행렬로 변환."""
    rows = []
    row = []
    cur = []
    in_quotes = False
    i = 0
    n = len(text)
    while i < n:
        ch = text[i]
        if in_quotes:
            if ch == '"':
                if i + 1 < n and text[i + 1] == '"':
                    cur.append('"')
                    i += 1
                else:
                    in_quotes = False
            else:
                cur.append(ch)
        elif ch == '"' and not cur:
            in_quotes = True
        elif ch == '\t':
            row.append(''.join(cur))
            cur = []
        elif ch == '\n':
            row.append(''.join(cur))
            cur = []
            rows.append(row)
            row = []
        elif ch != '\r':
            cur.append(ch)
        i += 1
    if cur or row:
        row.append(''.join(cur))
        rows.append(row)
    return rows


def _parse_budget_rows_no_header(rows):
    """헤더 없이 데이터만 복사된 경우: A=원가통계비목, B=산출내역,
    C=예산현액, E=원인행위금액(B) 위치 기준으로 파싱."""
    categories = []
    current = None

    def cell(row, idx):
        return row[idx] if idx < len(row) else None

    for row in rows:
        cost = _cell_text(cell(row, 0))
        desc = _cell_text(cell(row, 1))
        budget = _to_number(cell(row, 2))
        spent = _to_number(cell(row, 4)) or 0
        if not cost and not desc and budget is None:
            continue
        if any(key in cost for key in ('합계', '총계', '소계')):
            continue
        if not desc and budget is not None:
            current = {'name': cost or '미분류', 'items': []}
            categories.append(current)
            continue
        if desc:
            if current is None:
                current = {'name': '미분류', 'items': []}
                categories.append(current)
            current['items'].append({
                'cost_type': cost or '일반',
                'description': desc,
                'budget': budget or 0,
                'spent': spent,
            })
    if not categories or not any(c['items'] for c in categories):
        return None
    return categories


def parse_budget_clipboard(text):
    """클립보드 텍스트(엑셀 셀 복사) → [{'name': 세부항목, 'items': [...]}]."""
    text = (text or '').replace('\ufeff', '')
    if '\t' not in text:
        raise ValueError('탭으로 구분된 표 형식이 아닙니다.\n'
                         '엑셀에서 셀 범위(A~E열)를 복사한 뒤 다시 시도해주세요.')
    rows = _parse_tsv(text)
    categories = _parse_budget_sheet(rows) or _parse_budget_rows_no_header(rows)
    if not categories or not any(c['items'] for c in categories):
        raise ValueError('예산 데이터를 찾지 못했습니다.\n'
                         '세부항목 소계 행 포함, A열~E열 범위를 복사했는지 확인해주세요.')
    return categories


# ── 표 열 폭 / 높이 자동 맞춤 ──────────────────────────────────────
def _fit_columns(table):
    """표 폭에 맞춰 열 폭을 배분한다(넘치면 최소 폭까지 비례 축소).

    가로 스크롤바가 생기면 마지막 행이 잘리므로, 폭이 모자라면 열을 줄여
    스크롤바 없이 표 전체가 보이게 한다.
    """
    hdr = table.horizontalHeader()
    avail = table.viewport().width()
    n = len(_COL_W)
    if avail <= 0 or hdr.count() < n:
        return
    total = sum(_COL_W)
    if avail >= total:
        widths = list(_COL_W)
        widths[1] += avail - total            # 남는 폭은 산출내역 열이 흡수
    else:
        slack = [max(0, _COL_W[i] - _COL_MIN[i]) for i in range(n)]
        pool = sum(slack)
        if pool <= 0:
            return
        cut = min(total - avail, pool)
        widths = [max(_COL_MIN[i], int(_COL_W[i] - cut * slack[i] / pool))
                  for i in range(n)]
        over = sum(widths) - avail            # 반올림 초과분은 여유 있는 열에서 깎는다
        while over > 0:
            i = max(range(n), key=lambda k: widths[k] - _COL_MIN[k])
            if widths[i] <= _COL_MIN[i]:
                break
            widths[i] -= 1
            over -= 1
    for i, wd in enumerate(widths):
        if hdr.sectionSize(i) != wd:
            hdr.resizeSection(i, wd)


class _BudgetTable(QTableWidget):
    """폭에 맞춰 열 폭을 줄이고, 높이를 (헤더 + 행 높이 합)에 맞추는 표.

    기존 고정 높이(_HEADER_H + _ROW_H*n)는 실제 헤더 높이(46px)보다 작아
    마지막 행이 늘 조금씩 잘렸고, 가로 스크롤바가 뜨면 그만큼 더 잘렸다.
    """

    def sync_geometry(self):
        _fit_columns(self)
        self._sync_height()

    def _sync_height(self):
        hdr = self.horizontalHeader()
        rows_h = sum(self.rowHeight(r) for r in range(self.rowCount()))
        cols = sum(hdr.sectionSize(c) for c in range(hdr.count()))
        hbar_h = 0
        if cols > self.viewport().width() > 0:   # 열이 다 안 들어가면 스크롤바 높이 확보
            hbar_h = self.horizontalScrollBar().sizeHint().height()
        target = hdr.height() + rows_h + hbar_h
        if self.height() != target:
            self.setFixedHeight(target)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self.sync_geometry()


class BudgetWindow(QMainWindow):
    def __init__(self, parent=None):
        super().__init__(parent)
        _load_pretendard()
        self.categories = []
        self.items = []
        self.active_tab = 'ALL'
        self._items_by_id = {}
        self._row_refs = {}
        self._cat_refs = {}
        self._restore_window_state()
        self._init_ui()
        self.refresh()

    # ── UI 구성 ────────────────────────────────────────────────────
    def _init_ui(self):
        self.setWindowTitle('예산 관리')
        self.setMinimumSize(1000, 560)
        self.setStyleSheet(STYLE)
        self.setAcceptDrops(True)

        central = QWidget()
        self.setCentralWidget(central)
        root = QVBoxLayout(central)
        root.setContentsMargins(20, 16, 20, 16)
        root.setSpacing(14)

        # 헤더
        header = QHBoxLayout()
        header.setSpacing(10)
        icon_lbl = QLabel()
        icon_lbl.setPixmap(qta.icon('fa5s.calculator', color='#4f46e5').pixmap(26, 26))
        title_box = QVBoxLayout()
        title_box.setSpacing(0)
        lbl_title = QLabel('예산 관리')
        lbl_title.setStyleSheet('font-size: 17pt; font-weight: 800; color: #1e1b4b;')
        lbl_sub = QLabel('추경 세부항목 변경 대응 · 실시간 예산 잔액 및 집행률 · 항목 완료 기능')
        lbl_sub.setStyleSheet('font-size: 9pt; color: #64748b;')
        title_box.addWidget(lbl_title)
        title_box.addWidget(lbl_sub)
        header.addWidget(icon_lbl)
        header.addLayout(title_box)
        header.addStretch()

        btn_categories = QPushButton('예산 등록·관리')
        btn_categories.setObjectName('Dark')
        btn_categories.setIcon(qta.icon('fa5s.folder-open', color='white'))
        btn_categories.setIconSize(QSize(14, 14))
        btn_categories.clicked.connect(self._open_category_manager)
        header.addWidget(btn_categories)

        btn_reset = QPushButton('초기화')
        btn_reset.setObjectName('Danger')
        btn_reset.setIcon(qta.icon('fa5s.undo', color='white'))
        btn_reset.setIconSize(QSize(14, 14))
        btn_reset.clicked.connect(self._reset_all)
        header.addWidget(btn_reset)
        root.addLayout(header)

        # 요약 카드 3개
        cards = QHBoxLayout()
        cards.setSpacing(14)
        cards.addWidget(self._make_card('현재 편성 예산합 (예산현액)', 'CardValueIndigo',
                                        '세부항목별 예산현액 총합', 'allocated'))
        cards.addWidget(self._make_card('총 집행액 (지출)', 'CardValueAmber',
                                        '현재 편성 예산합 대비 집행률', 'spent'))
        cards.addWidget(self._make_card('총 예산 잔액', 'CardValueEmerald',
                                        '현재 편성 예산합 - 총 집행액', 'remaining'))
        root.addLayout(cards)

        # 탭 + 항목 추가 버튼
        toolbar = QFrame()
        toolbar.setObjectName('Toolbar')
        tb_layout = QHBoxLayout(toolbar)
        tb_layout.setContentsMargins(12, 10, 12, 10)
        tb_layout.setSpacing(8)
        self.tabs_container = QWidget()
        self.tabs_layout = QHBoxLayout(self.tabs_container)
        self.tabs_layout.setContentsMargins(0, 0, 0, 0)
        self.tabs_layout.setSpacing(8)
        tb_layout.addWidget(self.tabs_container)
        tb_layout.addStretch()
        self.btn_add_item = QPushButton('새 예산 항목 추가')
        self.btn_add_item.setObjectName('Primary')
        self.btn_add_item.setIcon(qta.icon('fa5s.plus-circle', color='white'))
        self.btn_add_item.setIconSize(QSize(14, 14))
        self.btn_add_item.clicked.connect(self._open_add_item)
        tb_layout.addWidget(self.btn_add_item)
        root.addWidget(toolbar)

        # 섹션 스크롤 영역
        self.scroll = QScrollArea()
        self.scroll.setWidgetResizable(True)
        self.sections_container = QWidget()
        self.sections_layout = QVBoxLayout(self.sections_container)
        self.sections_layout.setContentsMargins(0, 0, 6, 0)
        self.sections_layout.setSpacing(16)
        self.scroll.setWidget(self.sections_container)
        root.addWidget(self.scroll, stretch=1)

    def _make_card(self, title, value_obj, sub_text, key):
        card = QFrame()
        card.setObjectName('Card')
        v = QVBoxLayout(card)
        v.setContentsMargins(0, 0, 0, 0)
        v.setSpacing(0)

        body = QWidget()
        bl = QVBoxLayout(body)
        bl.setContentsMargins(16, 14, 16, 12)
        bl.setSpacing(4)

        lbl_title = QLabel(title)
        lbl_title.setObjectName('CardTitle')
        bl.addWidget(lbl_title)

        lbl_value = QLabel('0원')
        lbl_value.setObjectName(value_obj)
        if key == 'spent':
            row = QHBoxLayout()
            row.setSpacing(6)
            row.addWidget(lbl_value)
            row.addStretch()
            self.lbl_percent = QLabel('0.0%')
            self.lbl_percent.setObjectName('PercentBig')
            row.addWidget(self.lbl_percent)
            bl.addLayout(row)
        else:
            bl.addWidget(lbl_value)

        lbl_sub = QLabel(sub_text)
        lbl_sub.setObjectName('CardSub')
        bl.addWidget(lbl_sub)

        if key == 'spent':
            self.bar_spent = QProgressBar()
            self.bar_spent.setObjectName('SummaryBar')
            self.bar_spent.setRange(0, 100)
            self.bar_spent.setValue(0)
            self.bar_spent.setTextVisible(False)
            self.bar_spent.setFixedHeight(8)
            bl.addWidget(self.bar_spent)

        v.addWidget(body)

        if key == 'allocated':
            self.lbl_allocated = lbl_value
        elif key == 'remaining':
            self.lbl_remaining = lbl_value
        else:
            self.lbl_spent = lbl_value
        return card

    # ── 데이터 새로고침 ────────────────────────────────────────────
    def refresh(self, keep_scroll=True):
        pos = self.scroll.verticalScrollBar().value() if keep_scroll else 0
        self.categories = get_budget_categories()
        self.items = get_budget_items()
        self._items_by_id = {it['id']: it for it in self.items}
        valid_ids = {c['id'] for c in self.categories}
        if self.active_tab != 'ALL' and self.active_tab not in valid_ids:
            self.active_tab = 'ALL'
        self._rebuild_tabs()
        self._rebuild_sections()
        self._update_summary()
        self.btn_add_item.setEnabled(bool(self.categories))
        if keep_scroll:
            self.scroll.verticalScrollBar().setValue(pos)

    def _rebuild_tabs(self):
        while self.tabs_layout.count():
            w = self.tabs_layout.takeAt(0).widget()
            if w:
                w.setParent(None)
                w.deleteLater()

        all_btn = QPushButton('전체 보기')
        all_btn.setObjectName('TabActive' if self.active_tab == 'ALL' else 'TabBtn')
        all_btn.setCursor(Qt.PointingHandCursor)
        all_btn.clicked.connect(lambda: self._set_tab('ALL'))
        self.tabs_layout.addWidget(all_btn)

        for cat in self.categories:
            btn = QPushButton(cat['name'])
            btn.setObjectName('TabActive' if self.active_tab == cat['id'] else 'TabBtn')
            btn.setCursor(Qt.PointingHandCursor)
            btn.clicked.connect(lambda _, cid=cat['id']: self._set_tab(cid))
            self.tabs_layout.addWidget(btn)

    def _set_tab(self, tab):
        self.active_tab = tab
        self.refresh(keep_scroll=False)

    def _clear_sections(self):
        while self.sections_layout.count():
            w = self.sections_layout.takeAt(0).widget()
            if w:
                w.setParent(None)
                w.deleteLater()

    def _rebuild_sections(self):
        self._row_refs = {}
        self._cat_refs = {}
        self._clear_sections()

        if not self.categories:
            lbl = QLabel('등록된 세부항목이 없습니다.\n'
                         '[예산 등록·관리]에서 직접 추가하거나, '
                         '사업관리카드 엑셀 파일(.xls/.xlsx)을 창에 끌어다 놓으세요.\n'
                         '(엑셀에서 셀 범위를 복사해 Ctrl+V 붙여넣기도 가능합니다.)')
            lbl.setObjectName('EmptyBox')
            lbl.setAlignment(Qt.AlignCenter)
            self.sections_layout.addWidget(lbl)
            self.sections_layout.addStretch()
            return

        target = self.categories if self.active_tab == 'ALL' else \
            [c for c in self.categories if c['id'] == self.active_tab]

        for cat in target:
            cat_items = [it for it in self.items if it['category_id'] == cat['id']]
            self.sections_layout.addWidget(self._make_section(cat, cat_items))
        self.sections_layout.addStretch()

    def _make_section(self, cat, cat_items):
        section = QFrame()
        section.setObjectName('Section')
        v = QVBoxLayout(section)
        v.setContentsMargins(0, 0, 0, 0)
        v.setSpacing(0)

        # 헤더
        header = QFrame()
        header.setObjectName('SectionHeader')
        hl = QHBoxLayout(header)
        hl.setContentsMargins(16, 12, 16, 12)
        hl.setSpacing(8)

        icon_lbl = QLabel()
        icon_lbl.setPixmap(qta.icon('fa5s.folder', color='#a5b4fc').pixmap(16, 16))
        lbl_name = QLabel(cat['name'])
        lbl_name.setObjectName('SectionTitle')
        hl.addWidget(icon_lbl)
        hl.addWidget(lbl_name)
        hl.addStretch()

        chips = {}
        for key, prefix in (('budget', '예산합'), ('spent', '집행액'),
                            ('remaining', '잔액'), ('percent', '집행률')):
            chip = QLabel('')
            chip.setObjectName('Chip')
            chip.setTextFormat(Qt.RichText)
            chips[key] = chip
            hl.addWidget(chip)
        self._cat_refs[cat['id']] = chips

        v.addWidget(header)

        # 테이블
        table = _BudgetTable()
        table.setColumnCount(8)
        table.setHorizontalHeaderLabels(['원가통계비목', '산출내역 (내용)', '예산현액',
                                         '집행액 (입력)', '예산잔액', '집행률',
                                         '완료여부', '관리'])
        table.verticalHeader().setVisible(False)
        table.verticalHeader().setDefaultSectionSize(_ROW_H)
        table.setSelectionMode(QAbstractItemView.NoSelection)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setFocusPolicy(Qt.NoFocus)
        table.setShowGrid(False)
        table.setFont(QFont('Pretendard', 10))
        hdr = table.horizontalHeader()
        hdr.setHighlightSections(False)
        hdr.setSectionsClickable(False)
        hdr.setMinimumSectionSize(40)      # 기본 71px 바닥 때문에 좁은 폭에서 넘쳤다
        for col in range(8):
            hdr.setSectionResizeMode(col, QHeaderView.Interactive)
        for col, wd in enumerate(_COL_W):
            table.setColumnWidth(col, wd)

        if cat_items:
            table.setRowCount(len(cat_items))
            for idx, item in enumerate(cat_items):
                self._add_item_row(table, idx, item)
        else:
            table.setRowCount(1)
            msg = QTableWidgetItem('해당 세부항목에 등록된 예산 내역이 없습니다.')
            msg.setFlags(Qt.ItemIsEnabled)
            msg.setTextAlignment(Qt.AlignCenter)
            msg.setForeground(QColor('#94a3b8'))
            table.setItem(0, 0, msg)
            table.setSpan(0, 0, 1, 8)
            table.setRowHeight(0, 60)

        table.sync_geometry()
        v.addWidget(table)
        self._update_cat_header(cat['id'])
        return section

    def _add_item_row(self, table, row, item):
        is_done = bool(item['completed'])
        bg = _BG_DONE if is_done else (_BG_ZEBRA if row % 2 == 1 else _BG_WHITE)
        rem = item['budget'] - item['spent']
        pct = min(_pct(item['spent'], item['budget']), 100.0)

        def txt_item(text, align, color=_COLOR_TEXT, strike=False, bold=True):
            it = QTableWidgetItem(text)
            it.setFlags(Qt.ItemIsEnabled)
            it.setTextAlignment(align)
            it.setBackground(bg)
            f = QFont(table.font())
            f.setBold(bold)
            if strike:
                f.setStrikeOut(True)
            it.setFont(f)
            it.setForeground(_COLOR_DONE_TEXT if is_done else color)
            return it

        # 원가통계비목
        table.setItem(row, 0, txt_item(item['cost_type'], Qt.AlignLeft | Qt.AlignVCenter,
                                       _COLOR_COST, strike=is_done, bold=True))
        # 산출내역
        table.setItem(row, 1, txt_item(item['description'], Qt.AlignLeft | Qt.AlignVCenter,
                                       _COLOR_TEXT, strike=is_done))
        # 예산현액
        table.setItem(row, 2, txt_item(f'{_fmt(item["budget"])}원',
                                       Qt.AlignRight | Qt.AlignVCenter, strike=is_done))

        # 집행액 (인라인 편집)
        spent_wrap = QWidget()
        sl = QHBoxLayout(spent_wrap)
        sl.setContentsMargins(6, 5, 6, 5)
        sl.setSpacing(4)
        spent_edit = QLineEdit(f'{_fmt(item["spent"])}')
        spent_edit.setObjectName('SpentInput')
        spent_edit.setAlignment(Qt.AlignRight)
        spent_edit.setFrame(True)
        spent_edit.editingFinished.connect(
            lambda item_id=item['id'], e=spent_edit: self._on_spent_edited(item_id, e))
        if is_done:
            spent_edit.setDisabled(True)
            spent_edit.setFont(_make_strike_font(spent_edit.font()))
        sl.addWidget(spent_edit, stretch=1)
        won_lbl = QLabel('원')
        won_lbl.setStyleSheet('color: #64748b; background: transparent;')
        sl.addWidget(won_lbl)
        for w in (spent_wrap, spent_edit, won_lbl):
            w.setToolTip('집행액을 입력한 뒤 Enter 또는 다른 곳 클릭 시 저장됩니다.')
        table.setItem(row, 3, txt_item('', Qt.AlignRight | Qt.AlignVCenter, bold=False))
        table.setCellWidget(row, 3, spent_wrap)

        # 예산잔액
        rem_item = txt_item(f'{_fmt(rem)}원', Qt.AlignRight | Qt.AlignVCenter)
        if rem < 0 and not is_done:
            rem_item.setForeground(_COLOR_NEG)
            rem_item.setBackground(_BG_NEG)
        table.setItem(row, 4, rem_item)

        # 집행률
        pct_wrap = QWidget()
        pl = QHBoxLayout(pct_wrap)
        pl.setContentsMargins(6, 5, 8, 5)
        pl.setSpacing(6)
        pl.addStretch()
        bar = QProgressBar()
        bar.setRange(0, 100)
        bar.setValue(int(round(pct)))
        bar.setTextVisible(False)
        bar.setFixedSize(64, 8)
        pl.addWidget(bar)
        pct_lbl = QLabel(f'{pct:.1f}%')
        pct_lbl.setFixedWidth(54)
        pct_lbl.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        if is_done:
            bar.setStyleSheet('QProgressBar { background: #e2e8f0; border-radius: 4px; }'
                              'QProgressBar::chunk { background: #94a3b8; border-radius: 4px; }')
            pct_lbl.setStyleSheet('color: #64748b; background: transparent;')
            pct_lbl.setFont(_make_strike_font(pct_lbl.font()))
        else:
            pct_lbl.setStyleSheet('color: #1d4ed8; font-weight: 800; background: transparent;')
        pl.addWidget(pct_lbl)
        table.setItem(row, 5, txt_item('', Qt.AlignRight | Qt.AlignVCenter, bold=False))
        table.setCellWidget(row, 5, pct_wrap)

        # 완료 토글
        done_wrap = QWidget()
        dl = QHBoxLayout(done_wrap)
        dl.setContentsMargins(4, 6, 4, 6)
        dl.setSpacing(0)
        btn_done = QPushButton('완료 취소' if is_done else '완료')
        btn_done.setObjectName('DoneBtn' if is_done else 'CompleteBtn')
        btn_done.setCursor(Qt.PointingHandCursor)
        btn_done.setToolTip('완료를 취소합니다.' if is_done else '이 항목을 완료 처리합니다.')
        btn_done.clicked.connect(lambda _, iid=item['id']: self._toggle_complete(iid))
        dl.addWidget(btn_done)
        table.setItem(row, 6, txt_item('', Qt.AlignCenter, bold=False))
        table.setCellWidget(row, 6, done_wrap)

        # 수정/삭제
        act_wrap = QWidget()
        al = QHBoxLayout(act_wrap)
        al.setContentsMargins(2, 4, 2, 4)
        al.setSpacing(2)
        btn_edit = _icon_btn('fa5s.edit', '수정')
        btn_edit.clicked.connect(lambda _, iid=item['id']: self._open_edit_item(iid))
        btn_del = _icon_btn('fa5s.trash-alt', '삭제', color='#e11d48')
        btn_del.clicked.connect(lambda _, iid=item['id']: self._delete_item(iid))
        if is_done:
            btn_edit.setDisabled(True)
            btn_del.setDisabled(True)
        al.addWidget(btn_edit)
        al.addWidget(btn_del)
        table.setItem(row, 7, txt_item('', Qt.AlignCenter, bold=False))
        table.setCellWidget(row, 7, act_wrap)

        self._row_refs[item['id']] = {
            'remaining': rem_item,
            'pct_bar': bar,
            'pct_lbl': pct_lbl,
            'spent_edit': spent_edit,
            'bg': bg,
        }

    # ── 실시간 재계산 ──────────────────────────────────────────────
    def _on_spent_edited(self, item_id, line_edit):
        item = self._items_by_id.get(item_id)
        if not item or item['completed']:
            return
        val = _parse_amount(line_edit.text())
        line_edit.setText(f'{val:,}')
        if val == item['spent']:
            return
        set_budget_item_spent(item_id, val)
        item['spent'] = val
        self._recalc_item(item_id)
        self._recalc_cat_header(item['category_id'])
        self._update_summary()

    def _recalc_item(self, item_id):
        refs = self._row_refs.get(item_id)
        item = self._items_by_id.get(item_id)
        if not refs or not item:
            return
        rem = item['budget'] - item['spent']
        pct = min(_pct(item['spent'], item['budget']), 100.0)
        refs['remaining'].setText(f'{_fmt(rem)}원')
        if rem < 0:
            refs['remaining'].setForeground(_COLOR_NEG)
            refs['remaining'].setBackground(_BG_NEG)
        else:
            refs['remaining'].setForeground(_COLOR_TEXT)
            refs['remaining'].setBackground(refs['bg'])
        refs['pct_bar'].setValue(int(round(pct)))
        refs['pct_lbl'].setText(f'{pct:.1f}%')

    # ── 합계 계산 ──────────────────────────────────────────────────
    def _cat_totals(self, cat_id):
        budget = spent = 0
        for it in self.items:
            if it['category_id'] == cat_id:
                budget += it['budget']
                spent += it['spent']
        return budget, spent

    def _update_cat_header(self, cat_id):
        chips = self._cat_refs.get(cat_id)
        if not chips:
            return
        budget, spent = self._cat_totals(cat_id)
        rem = budget - spent
        pct = _pct(spent, budget)
        chips['budget'].setText(f'<span style="color:#ffffff">예산합: </span>'
                                f'<b style="color:#ffffff">{_fmt(budget)}원</b>')
        chips['spent'].setText(f'<span style="color:#ffffff">집행액: </span>'
                               f'<b style="color:#fcd34d">{_fmt(spent)}원</b>')
        chips['remaining'].setText(f'<span style="color:#ffffff">잔액: </span>'
                                   f'<b style="color:#6ee7b7">{_fmt(rem)}원</b>')
        chips['percent'].setText(f'<span style="color:#ffffff">집행률: </span>'
                                 f'<b style="color:#67e8f9">{pct:.1f}%</b>')

    def _recalc_cat_header(self, cat_id):
        self._update_cat_header(cat_id)

    def _update_summary(self):
        total_allocated = sum(it['budget'] for it in self.items)
        total_spent = sum(it['spent'] for it in self.items)
        total_remaining = total_allocated - total_spent
        pct = _pct(total_spent, total_allocated)
        self.lbl_allocated.setText(f'{_fmt(total_allocated)}원')
        self.lbl_spent.setText(f'{_fmt(total_spent)}원')
        self.lbl_remaining.setText(f'{_fmt(total_remaining)}원')
        self.lbl_percent.setText(f'{pct:.1f}%')
        self.bar_spent.setValue(int(round(min(pct, 100.0))))

    # ── 동작: 항목/세부항목 ────────────────────────────────────────
    def _toggle_complete(self, item_id):
        toggle_budget_item_completed(item_id)
        self.refresh()

    def _open_add_item(self):
        if not self.categories:
            QMessageBox.information(self, '세부항목 필요',
                                    '먼저 [예산 등록·관리]에서 세부항목을 추가하거나 '
                                    '사업관리카드 엑셀 파일을 창에 끌어다 놓으세요.')
            self._open_category_manager()
            return
        default_cat = self.active_tab if self.active_tab != 'ALL' else self.categories[0]['id']
        dlg = _ItemDialog(self.categories, None, default_cat, self)
        if dlg.exec_() == QDialog.Accepted and dlg.result_data:
            d = dlg.result_data
            add_budget_item(d['category_id'], d['cost_type'], d['description'],
                            d['budget'], d['spent'])
            self.refresh(keep_scroll=False)

    def _open_edit_item(self, item_id):
        item = self._items_by_id.get(item_id)
        if not item:
            return
        dlg = _ItemDialog(self.categories, item, item['category_id'], self)
        if dlg.exec_() == QDialog.Accepted and dlg.result_data:
            d = dlg.result_data
            update_budget_item(item_id, d['category_id'], d['cost_type'],
                               d['description'], d['budget'], d['spent'])
            self.refresh()

    def _delete_item(self, item_id):
        item = self._items_by_id.get(item_id)
        if not item:
            return
        ans = QMessageBox.question(self, '항목 삭제',
                                   f"'{item['description']}' 항목을 삭제하시겠습니까?",
                                   QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
        if ans == QMessageBox.Yes:
            delete_budget_item(item_id)
            self.refresh()

    def _open_category_manager(self):
        _CategoryManagerDialog(self, self.refresh).exec_()

    def _reset_all(self):
        if not self.categories and not self.items:
            QMessageBox.information(self, '초기화', '삭제할 데이터가 없습니다.')
            return
        ans = QMessageBox.question(self, '초기화',
                                   '모든 세부항목과 예산 항목이 삭제됩니다. 계속하시겠습니까?',
                                   QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
        if ans == QMessageBox.Yes:
            clear_budget_data()
            self.active_tab = 'ALL'
            self.refresh(keep_scroll=False)

    # ── 엑셀 가져오기 / 클립보드 붙여넣기 ──────────────────────────
    def _choose_excel(self, parent=None):
        parent = parent or self
        path, _ = QFileDialog.getOpenFileName(
            parent, '예산 엑셀 파일 선택', '',
            '엑셀 파일 (*.xls *.xlsx)')
        if path:
            self._import_excel(path, parent=parent)

    def _import_excel(self, path, parent=None):
        try:
            parsed = parse_budget_excel(path)
        except Exception as e:
            QMessageBox.critical(parent or self, '가져오기 실패',
                                 f'엑셀 파일을 읽을 수 없습니다.\n\n{e}')
            return
        self._import_parsed(parsed, f'파일: {os.path.basename(path)}', parent=parent)

    def _paste_from_clipboard(self):
        text = QApplication.clipboard().text()
        if not text.strip():
            QMessageBox.information(self, '붙여넣기',
                                    '클립보드가 비어 있습니다.\n'
                                    '엑셀에서 셀 범위를 복사한 뒤 다시 시도해주세요.')
            return
        try:
            parsed = parse_budget_clipboard(text)
        except Exception as e:
            QMessageBox.warning(self, '붙여넣기 실패', str(e))
            return
        self._import_parsed(parsed, '출처: 클립보드 붙여넣기')

    def _import_parsed(self, parsed, source_text, parent=None):
        parent = parent or self
        if not parsed or not any(c['items'] for c in parsed):
            QMessageBox.warning(parent, '가져오기 실패',
                                '예산 데이터를 찾지 못했습니다.\n'
                                "사업관리카드(예산) 형식('예산현액', '산출내역' 열)인지 확인해주세요.")
            return
        has_existing = bool(self.categories or self.items)
        dlg = _ImportPreviewDialog(source_text, parsed, has_existing, parent)
        if dlg.exec_() != QDialog.Accepted:
            return
        replace = (dlg.mode == 'replace')
        self._apply_import(parsed, replace)
        n_cat = len(parsed)
        n_item = sum(len(c['items']) for c in parsed)
        action = '교체' if replace else '추가'
        QMessageBox.information(
            parent, '가져오기 완료',
            f'세부항목 {n_cat}개, 예산 항목 {n_item}개를 {action}했습니다.')

    def _apply_import(self, categories, replace):
        if replace:
            clear_budget_data()
        existing = {c['name']: c['id'] for c in get_budget_categories()}
        for cat in categories:
            cid = existing.get(cat['name'])
            if cid is None:
                cid = add_budget_category(cat['name'])
                if cid is None:
                    existing = {c['name']: c['id'] for c in get_budget_categories()}
                    cid = existing.get(cat['name'])
                if cid is None:
                    continue
                existing[cat['name']] = cid
            for it in cat['items']:
                add_budget_item(cid, it['cost_type'], it['description'],
                                it['budget'], it['spent'])
        self.active_tab = 'ALL'
        self.refresh(keep_scroll=False)

    # ── 드래그 & 드롭 ─────────────────────────────────────────────
    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            for url in event.mimeData().urls():
                if url.toLocalFile().lower().endswith(_EXCEL_EXTS):
                    event.acceptProposedAction()
                    return
        event.ignore()

    def dragMoveEvent(self, event):
        self.dragEnterEvent(event)

    def dropEvent(self, event):
        for url in event.mimeData().urls():
            path = url.toLocalFile()
            if path.lower().endswith(_EXCEL_EXTS):
                event.acceptProposedAction()
                self._import_excel(path)
                return
        event.ignore()

    # ── Ctrl+V 붙여넣기 ────────────────────────────────────────────
    def keyPressEvent(self, event):
        if event.matches(QKeySequence.Paste) and self._can_paste_import():
            self._paste_from_clipboard()
            return
        super().keyPressEvent(event)

    def _can_paste_import(self):
        w = QApplication.focusWidget()
        return not isinstance(w, (QLineEdit, QTextEdit, QPlainTextEdit, QComboBox))

    # ── 창 위치/크기 저장·복원 ────────────────────────────────────
    def _restore_window_state(self):
        try:
            with open(_STATE_FILE, 'r', encoding='utf-8') as f:
                s = json.load(f)
            self.setGeometry(s['x'], s['y'], s['width'], s['height'])
        except Exception:
            self.setGeometry(150, 150, 1120, 720)

    def showEvent(self, event):
        super().showEvent(event)
        self.refresh(keep_scroll=False)

    def closeEvent(self, event):
        try:
            os.makedirs(os.path.dirname(_STATE_FILE), exist_ok=True)
            g = self.geometry()
            with open(_STATE_FILE, 'w', encoding='utf-8') as f:
                json.dump({'x': g.x(), 'y': g.y(),
                           'width': g.width(), 'height': g.height()}, f)
        except Exception:
            pass
        super().closeEvent(event)


# ── 엑셀 가져오기 미리보기 다이얼로그 ────────────────────────────
class _ImportPreviewDialog(QDialog):
    def __init__(self, source_text, categories, has_existing, parent=None):
        super().__init__(parent)
        self.setWindowTitle('엑셀 가져오기 미리보기')
        self.setStyleSheet(STYLE)
        self.setMinimumSize(760, 560)
        self.mode = None

        total_budget = sum(it['budget'] for c in categories for it in c['items'])
        total_spent = sum(it['spent'] for c in categories for it in c['items'])
        n_items = sum(len(c['items']) for c in categories)

        v = QVBoxLayout(self)
        v.setContentsMargins(20, 18, 20, 18)
        v.setSpacing(10)

        lbl_file = QLabel(source_text)
        lbl_file.setStyleSheet('font-size: 10pt; font-weight: 700; color: #1e1b4b;')
        v.addWidget(lbl_file)

        lbl_sum = QLabel(
            f"세부항목 <b>{len(categories)}</b>개 · 예산 항목 <b>{n_items}</b>개 · "
            f"예산합계 <b>{_fmt(total_budget)}원</b> · 집행합계 <b>{_fmt(total_spent)}원</b> · "
            f"잔액 <b>{_fmt(total_budget - total_spent)}원</b>")
        lbl_sum.setTextFormat(Qt.RichText)
        lbl_sum.setStyleSheet('color: #475569; font-size: 9.5pt;')
        v.addWidget(lbl_sum)

        table = QTableWidget()
        table.setColumnCount(5)
        table.setHorizontalHeaderLabels(['구분', '원가통계비목', '산출내역 (내용)',
                                         '예산현액', '집행액'])
        table.verticalHeader().setVisible(False)
        table.setSelectionMode(QAbstractItemView.NoSelection)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setFocusPolicy(Qt.NoFocus)
        table.setShowGrid(False)
        table.setFont(QFont('Pretendard', 9))
        hdr = table.horizontalHeader()
        hdr.setHighlightSections(False)
        hdr.setSectionsClickable(False)
        for col in range(5):
            hdr.setSectionResizeMode(col, QHeaderView.Interactive)
        hdr.setSectionResizeMode(2, QHeaderView.Stretch)
        table.setColumnWidth(0, 70)
        table.setColumnWidth(1, 150)
        table.setColumnWidth(3, 110)
        table.setColumnWidth(4, 110)
        table.setStyleSheet(
            'QTableWidget { background: white; border: 1px solid #e2e8f0; border-radius: 8px; }'
            'QTableWidget::item { border-bottom: 1px solid #f1f5f9; padding: 3px 6px; }'
            'QHeaderView::section { background: #e2e8f0; color: #334155; padding: 5px;'
            ' border: none; border-right: 1px solid #cbd5e1; font-weight: 700; }')

        rows = []
        for cat in categories:
            cat_budget = sum(it['budget'] for it in cat['items'])
            cat_spent = sum(it['spent'] for it in cat['items'])
            rows.append(('cat', cat['name'], f"항목 {len(cat['items'])}개", cat_budget, cat_spent))
            for it in cat['items']:
                rows.append(('item', it['cost_type'], it['description'],
                             it['budget'], it['spent']))

        table.setRowCount(len(rows))
        for r, (kind, c0, c1, c2, c3) in enumerate(rows):
            texts = (['[세부항목]', c0, c1, f'{_fmt(c2)}원', f'{_fmt(c3)}원']
                     if kind == 'cat' else
                     ['', c0, c1, f'{_fmt(c2)}원', f'{_fmt(c3)}원'])
            for col, text in enumerate(texts):
                cell = QTableWidgetItem(text)
                cell.setFlags(Qt.ItemIsEnabled)
                if col >= 3:
                    cell.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                f = QFont(table.font())
                if kind == 'cat':
                    f.setBold(True)
                    cell.setBackground(QColor('#eef2ff'))
                    cell.setForeground(QColor('#3730a3'))
                cell.setFont(f)
                table.setItem(r, col, cell)
            table.setRowHeight(r, 26)
        v.addWidget(table, stretch=1)

        mode_row = QVBoxLayout()
        mode_row.setSpacing(4)
        self.cb_replace = QCheckBox('기존 데이터를 모두 지우고 교체')
        self.cb_append = QCheckBox('기존 데이터에 추가')
        self._mode_group = QButtonGroup(self)
        self._mode_group.setExclusive(True)
        self._mode_group.addButton(self.cb_replace)
        self._mode_group.addButton(self.cb_append)
        mode_row.addWidget(self.cb_replace)
        mode_row.addWidget(self.cb_append)
        if not has_existing:
            self.cb_append.setEnabled(False)
        v.addLayout(mode_row)

        btns = QHBoxLayout()
        btns.addStretch()
        btn_cancel = QPushButton('취소')
        btn_cancel.setObjectName('Ghost')
        btn_cancel.clicked.connect(self.reject)
        btn_import = QPushButton('가져오기')
        btn_import.setObjectName('Primary')
        btn_import.clicked.connect(self._accept)
        btns.addWidget(btn_cancel)
        btns.addWidget(btn_import)
        v.addLayout(btns)

    def _accept(self):
        if self.cb_append.isChecked():
            self.mode = 'append'
        elif self.cb_replace.isChecked():
            self.mode = 'replace'
        else:
            QMessageBox.warning(self, '가져오기 방식 선택',
                                '기존 데이터를 교체할지 추가할지 먼저 선택해주세요.')
            return
        self.accept()


# ── 예산 항목 추가/수정 다이얼로그 ────────────────────────────────
class _ItemDialog(QDialog):
    def __init__(self, categories, item, default_cat_id, parent=None):
        super().__init__(parent)
        self.setWindowTitle('예산 항목 수정' if item else '새 예산 항목 추가')
        self.setStyleSheet(STYLE)
        self.setMinimumWidth(420)
        self.result_data = None
        self._item = item

        v = QVBoxLayout(self)
        v.setContentsMargins(20, 18, 20, 18)
        v.setSpacing(10)

        v.addWidget(self._label('세부항목'))
        self.combo_cat = QComboBox()
        for cat in categories:
            self.combo_cat.addItem(cat['name'], cat['id'])
        idx = self.combo_cat.findData(default_cat_id)
        if idx >= 0:
            self.combo_cat.setCurrentIndex(idx)
        v.addWidget(self.combo_cat)

        v.addWidget(self._label('원가통계비목'))
        self.edit_cost = QLineEdit()
        self.edit_cost.setPlaceholderText('예: 교육운영비, 운영수당')
        v.addWidget(self.edit_cost)

        v.addWidget(self._label('산출내역 (내용)'))
        self.edit_desc = QLineEdit()
        self.edit_desc.setPlaceholderText('예: 교과융합캠프 간식비')
        v.addWidget(self.edit_desc)

        amount_row = QHBoxLayout()
        amount_row.setSpacing(12)
        budget_col = QVBoxLayout()
        budget_col.setSpacing(4)
        budget_col.addWidget(self._label('예산현액 (원)'))
        self.edit_budget = QLineEdit('0')
        self.edit_budget.setAlignment(Qt.AlignRight)
        _attach_amount_format(self.edit_budget)
        budget_col.addWidget(self.edit_budget)
        spent_col = QVBoxLayout()
        spent_col.setSpacing(4)
        spent_col.addWidget(self._label('집행액 (지출액)'))
        self.edit_spent = QLineEdit('0')
        self.edit_spent.setAlignment(Qt.AlignRight)
        _attach_amount_format(self.edit_spent)
        spent_col.addWidget(self.edit_spent)
        amount_row.addLayout(budget_col)
        amount_row.addLayout(spent_col)
        v.addLayout(amount_row)

        btns = QHBoxLayout()
        btns.addStretch()
        btn_cancel = QPushButton('취소')
        btn_cancel.setObjectName('Ghost')
        btn_cancel.clicked.connect(self.reject)
        btn_save = QPushButton('저장')
        btn_save.setObjectName('Primary')
        btn_save.clicked.connect(self._save)
        btns.addWidget(btn_cancel)
        btns.addWidget(btn_save)
        v.addLayout(btns)

        if item:
            self.edit_cost.setText(item['cost_type'])
            self.edit_desc.setText(item['description'])
            self.edit_budget.setText(f'{item["budget"]:,}')
            self.edit_spent.setText(f'{item["spent"]:,}')
        self.edit_desc.setFocus()

    @staticmethod
    def _label(text):
        lbl = QLabel(text)
        lbl.setStyleSheet('font-size: 9.5pt; font-weight: 700; color: #475569;')
        return lbl

    def _save(self):
        desc = self.edit_desc.text().strip()
        if not desc:
            QMessageBox.warning(self, '저장 실패', '산출내역(내용)을 입력해주세요.')
            return
        cost = self.edit_cost.text().strip() or '일반'
        self.result_data = {
            'category_id': self.combo_cat.currentData(),
            'cost_type': cost,
            'description': desc,
            'budget': _parse_amount(self.edit_budget.text()),
            'spent': _parse_amount(self.edit_spent.text()),
        }
        self.accept()


# ── 세부항목(추경) 관리 다이얼로그 ────────────────────────────────
class _CategoryManagerDialog(QDialog):
    def __init__(self, parent, on_changed):
        super().__init__(parent)
        self.setWindowTitle('세부항목(추경) 관리')
        self.setStyleSheet(STYLE)
        self.setMinimumWidth(460)
        self.setMinimumHeight(380)
        self._on_changed = on_changed

        v = QVBoxLayout(self)
        v.setContentsMargins(20, 18, 20, 18)
        v.setSpacing(10)

        lbl = QLabel('추경 발생 시 세부항목을 새로 추가하거나 명칭 변경, 삭제가 가능합니다.')
        lbl.setStyleSheet('color: #64748b; font-size: 9pt;')
        v.addWidget(lbl)

        add_row = QHBoxLayout()
        add_row.setSpacing(8)
        self.edit_new = QLineEdit()
        self.edit_new.setPlaceholderText('새 세부항목 이름 (예: [목적]신규사업)')
        self.edit_new.returnPressed.connect(self._add_category)
        btn_add = QPushButton('추가')
        btn_add.setObjectName('Primary')
        btn_add.clicked.connect(self._add_category)
        add_row.addWidget(self.edit_new, stretch=1)
        add_row.addWidget(btn_add)
        v.addLayout(add_row)

        self.list_scroll = QScrollArea()
        self.list_scroll.setWidgetResizable(True)
        self.list_scroll.setStyleSheet(
            'QScrollArea { border: 1px solid #e2e8f0; border-radius: 8px; background: white; }')
        self.list_container = QWidget()
        self.list_layout = QVBoxLayout(self.list_container)
        self.list_layout.setContentsMargins(6, 6, 6, 6)
        self.list_layout.setSpacing(4)
        self.list_scroll.setWidget(self.list_container)
        v.addWidget(self.list_scroll, stretch=1)

        btns = QHBoxLayout()
        self.btn_import = QPushButton('엑셀 가져오기')
        self.btn_import.setObjectName('Ghost')
        self.btn_import.setIcon(qta.icon('fa5s.file-import', color='#334155'))
        self.btn_import.setIconSize(QSize(14, 14))
        self.btn_import.setToolTip('사업관리카드(예산) 엑셀 파일(.xls/.xlsx)을 가져와 '
                                   '세부항목을 자동 생성합니다.')
        self.btn_import.clicked.connect(lambda: self._import_excel())
        btns.addWidget(self.btn_import)
        btns.addStretch()
        btn_close = QPushButton('닫기')
        btn_close.setObjectName('Dark')
        btn_close.clicked.connect(self.accept)
        btns.addWidget(btn_close)
        v.addLayout(btns)

        self._rebuild_list()

    def _import_excel(self):
        win = self.parent()
        if win is None or not hasattr(win, '_choose_excel'):
            return
        win._choose_excel(parent=self)
        self._rebuild_list()

    def _rebuild_list(self):
        while self.list_layout.count():
            w = self.list_layout.takeAt(0).widget()
            if w:
                w.setParent(None)
                w.deleteLater()

        cats = get_budget_categories()
        if not cats:
            empty = QLabel('등록된 세부항목이 없습니다.\n'
                           '아래 [엑셀 가져오기]로 사업관리카드 파일을 불러오거나,\n'
                           '위에서 직접 추가하세요.')
            empty.setAlignment(Qt.AlignCenter)
            empty.setStyleSheet('color: #94a3b8; padding: 20px;')
            self.list_layout.addWidget(empty)
        for cat in cats:
            row = QWidget()
            rl = QHBoxLayout(row)
            rl.setContentsMargins(2, 2, 2, 2)
            rl.setSpacing(6)
            edit = QLineEdit(cat['name'])
            edit.editingFinished.connect(
                lambda cid=cat['id'], e=edit, old=cat['name']: self._rename_category(cid, e, old))
            btn_del = _icon_btn('fa5s.trash-alt', '세부항목 삭제', color='#e11d48')
            btn_del.clicked.connect(lambda _, cid=cat['id']: self._delete_category(cid))
            rl.addWidget(edit, stretch=1)
            rl.addWidget(btn_del)
            self.list_layout.addWidget(row)
        self.list_layout.addStretch()

    def _add_category(self):
        name = self.edit_new.text().strip()
        if not name:
            return
        if add_budget_category(name) is None:
            QMessageBox.warning(self, '추가 실패', '이미 존재하는 세부항목 이름입니다.')
            return
        self.edit_new.clear()
        self._rebuild_list()
        self._on_changed()

    def _rename_category(self, cat_id, edit, old_name):
        new_name = edit.text().strip()
        if not new_name or new_name == old_name:
            edit.setText(old_name)
            return
        if not rename_budget_category(cat_id, new_name):
            QMessageBox.warning(self, '변경 실패', '이미 존재하는 세부항목 이름입니다.')
            edit.setText(old_name)
            return
        self._on_changed()

    def _delete_category(self, cat_id):
        cats = {c['id']: c for c in get_budget_categories()}
        cat = cats.get(cat_id)
        if not cat:
            return
        count = len([it for it in get_budget_items(cat_id)])
        if count > 0:
            ans = QMessageBox.question(
                self, '세부항목 삭제',
                f"'{cat['name']}' 세부항목에 속한 예산 항목이 {count}개 있습니다.\n"
                '정말 함께 삭제하시겠습니까?',
                QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
            if ans != QMessageBox.Yes:
                return
        delete_budget_category(cat_id)
        self._rebuild_list()
        self._on_changed()


# ── 싱글턴 창 열기 ────────────────────────────────────────────────
_instance = None


def _on_destroyed(*_args):
    global _instance
    _instance = None


def open_budget_window():
    global _instance
    if _instance is None:
        _instance = BudgetWindow()
        _instance.destroyed.connect(_on_destroyed)
    _instance.show()
    _instance.raise_()
    _instance.activateWindow()
